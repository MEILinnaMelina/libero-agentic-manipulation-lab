# -*- coding: utf-8 -*-
"""Build the two-day report from retained project evidence. No experiments or API calls."""
import sys, json
from pathlib import Path
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'work/docdeps'))
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = ROOT / 'reports/work_report/LIBERO_工作报告_2026-09-07至08.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
d = Document()
sec = d.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21)
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.7)
sec.left_margin = sec.right_margin = Cm(2)
sec.header_distance = sec.footer_distance = Cm(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    s = d.styles[name]
    s.font.name = 'Microsoft YaHei'
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    s.font.color.rgb = RGBColor.from_string('172332')
d.styles['Normal'].font.size = Pt(10.5)
d.styles['Normal'].paragraph_format.line_spacing = 1.18
d.styles['Normal'].paragraph_format.space_after = Pt(7)
d.styles['Title'].font.size = Pt(28)
d.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
for border in d.styles['Title']._element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
d.styles['Subtitle'].font.italic = False
d.styles['Heading 1'].font.size = Pt(19)
d.styles['Heading 2'].font.size = Pt(13)
for name in ['Heading 1', 'Heading 2']:
    d.styles[name].paragraph_format.space_before = Pt(10)
    d.styles[name].paragraph_format.space_after = Pt(7)
d.styles['Caption'].font.size = Pt(9)
d.styles['Caption'].font.italic = False
h = sec.header.paragraphs[0]
h.text = 'LIBERO 机器人操作评测    /    两日工作记录'
h.style = d.styles['Caption']
f = sec.footer.paragraphs[0]
f.alignment = 2
f.add_run('2026 年 9 月 8 日    ·    ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); f._p.append(field)
for r in f.runs: r.font.size = Pt(8)
d.core_properties.title = 'LIBERO 机器人操作评测与固定语义修复工作报告'
d.core_properties.author = 'Codex'
d.core_properties.subject = '2026年9月7日至8日的实现、实验、故障修复与版本归档'

def p(s, style=None): return d.add_paragraph(s, style)
def head(s): d.add_heading(s, 2)
def page(n, title):
    d.add_page_break(); d.add_heading(f'{n:02d}  {title}', 1)
def table(headers, rows, widths=None):
    t = d.add_table(rows=1, cols=len(headers)); t.style = 'Light Shading Accent 1'
    t.autofit = False
    for i,x in enumerate(headers): t.rows[0].cells[i].text = x
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for i,x in enumerate(row): cells[i].text = str(x)
    for row in t.rows:
        nosplit = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(nosplit)
        for i,c in enumerate(row.cells):
            if widths: c.width = Cm(widths[i])
            for pp in c.paragraphs:
                pp.paragraph_format.space_after = Pt(4)
                pp.paragraph_format.space_before = Pt(4)
                pp.paragraph_format.line_spacing = 1.08
                for r in pp.runs: r.font.size = Pt(9)
    p('')
    return t
def evidence(s):
    pp=p('证据  '+s, 'Caption'); pp.paragraph_format.space_before=Pt(7)

p('2026 年 9 月 7 日至 8 日', 'Subtitle')
d.add_heading('LIBERO 机器人操作评测\n与固定语义修复工作报告', 0)
p('实现迁移 · 历史评测核查 · 八项任务修复 · 录像与版本归档', 'Subtitle')
p('项目  LIBERO-GPT6-Eval\n整理  Codex\n版本  v0.1.0-fixed-repair\n报告日期  2026 年 9 月 8 日')
head('本次工作的实际结论')
p('两天的工作已完成独立评测工程、Windows 仿真适配、GPT 与固定语义执行链路，以及原先失败的 8 个任务的逐项物理操作修复。最终统一版本在开发初始状态 0、1、2 上完成 24/24 次固定语义成功验证，并保存全部 24 段双视角录像。')
table(['本轮修复结果','API 使用','证据完整性'],[['8 个任务各 3/3 成功','修复实验 0 次调用\n输入输出均为 0 token','24 段录像逐一解码\n1 条长轨迹精确回放']],[5.7,5.6,5.7])
p('此前使用 API 的正式评测仍处于暂停状态：计划 470 回合，留存 83 个结果目录，其中成功 33、额度错误 47、中断 3，其余 387 回合未执行。该批次没有形成完整正式成绩。开发阶段也确实调用过 API；其他任务在开发阶段失败，主要源于抓取、几何与控制实现，不能全部归因于余额不足。')
p('本报告以本地代码、结果 JSON、调用元数据、修复日志与录像为依据。日期按本次会话的两日工作阶段归纳，不为缺少可靠时间戳的操作补写具体时刻。当前报告和 Git 归档没有启动新的仿真实验或模型调用。')
evidence('reports/FIXED_REPAIR_STATUS.json；reports/PAUSED_STATUS.json；reports/fixed_repair/summary.json')

page(1,'目标范围与两日工作进展')
head('目标和范围的演进')
p('原始目标是根据 TASK_LIST.md 将 RoboEval 的状态辅助代理方法迁移至本地 LIBERO-10，并建立可核验的评测流程。随着开发与正式运行暴露问题，工作先转向区分 API 成本、执行故障和额度中断；之后按用户指令暂停付费评测，仅排查固定语义下失败的任务 1、2、3、4、5、6、8、9，直到成功并录像。')
table(['阶段','完成的工作','形成的结果'],[
['9 月 7 日\n迁移与运行','源代码及资产快照、依赖锁定、Windows 修补、任务清单、控制与规划接口、开发实验及正式运行','独立工程与日志体系；识别内存、动作语义和技能问题；正式运行后因额度耗尽暂停'],
['9 月 8 日\n固定语义修复','禁用网络的诊断入口；共享坐标系修复；逐任务抓取和运动路径迭代；保留失败尝试','首次统一验证 23/24；修复清障边界后第二次统一验证 24/24'],
['本次交付\n工作整理','核对历史口径、制作 Word 报告、整理源码与实验资产、Git 版本和远端仓库归档','详细报告、可追溯版本、实验录像和原始证据']],[2.7,7.3,7])
head('任务清单完成度')
p('当前 TASK_LIST.md 标记 33/38 项完成，未关闭的是 T33、T34、T36、T37、T38。复选框表示对应工程验收证据，不代表全部机器人任务已经稳定解决。固定语义修复增加了充分的开发验证，但尚未补齐完整正式评测、匹配消融和正式研究报告所需的实验数据。')
head('本次明确交付的内容')
p('交付包括独立适配工程、保留来源和许可证的 vendor 快照、配置与测试、历史成功和失败记录、91 条固定语义修复尝试索引、最终 24 回合统一验证、录像核验信息、精确回放结果，以及本工作报告。原始 E:\\djf\\LIBERO 和 RoboEval-main 仓库没有写入本项目实验输出。')
evidence('TASK_LIST.md；docs/TASK_LIST.original.md；docs/CHANGELOG.md；docs/fixed_repair_attempts.jsonl')

page(2,'工程实现与运行环境')
p('工程采用本地锁定的 LIBERO 与 RoboEval 快照，新增适配层独立放在 src/libero_eval。运行依赖固定为 Python 3.8.13、NumPy 1.22.4、robosuite 1.4.0、MuJoCo 2.3.7、PyTorch 2.0.1 CPU、SciPy 1.10.1、Gym 0.25.2 与 BDDL 1.0.1。本阶段没有训练模型。')
table(['组件','实现内容及作用'],[
['env.py 与 scene.py','封装官方初始化、动作与终止；导出物体、关节、接触、持物和目标状态；统一世界坐标与四元数含义。'],
['motion.py','在独立的 MuJoCo 数据对象中做阻尼、多初值雅可比 IK；检查机器人和被持物碰撞；路径采样间隔不大于 4 cm。'],
['skills.py','抓取、抬升、搬运、释放、放置、开关、抽屉和门操作；通过正常控制接口执行并反馈前置条件与失败原因。'],
['planner.py','通过 HTTPS Responses 请求严格结构化计划；要求返回 gpt-6-astra，禁止其他模型回退；限制调用、输出和重规划预算。'],
['budget.py 与 runner','SQLite 预算预留及用量记录；保存配置、请求响应、决策、技能、动作状态、结果及 MP4；区分失败与中断。'],
['调试与报告脚本','固定语义诊断、来源验证、清单生成、版本冻结、分母统计、离线回放及报告生成。']],[3.8,13.2])
head('Windows 适配和内存问题')
p('修复 robosuite 日志写入 /tmp、MuJoCo DLL 搜索路径和渲染后端选择；在 Windows 使用 GLFW。最初四个仿真进程耗尽系统提交内存，该批开发结果被保留并排除。随后缩减 MuJoCo 工作区预留并限制 BLAS/OpenMP 为单线程，正式运行改为三个独立 worker。')
p('模型编译会重新计算实际 nstack，不能把配置中的 200 万直接理解为最终内存。任务 0 编译后约为 4089 万个 double。修改没有降低接触容量或控制器参数；对一条 262 动作轨迹做修改前后对照，状态逐项完全一致，作为内存调整保持物理执行一致的证据。')
evidence('configs/requirements-sim.lock；docs/provenance/windows_patches.json；reports/memory_equivalence.json；src/libero_eval/')

page(3,'实验协议与结果解释边界')
table(['项目','实际执行口径'],[
['任务和初始状态','LIBERO-10；每任务 50 个官方初始状态。0、1、2 用于开发，3 至 49 用于正式评测，共 470 回合。'],
['观测','使用物体几何、关节、接触与 BDDL 目标等特权状态。不是仅 RGB 的视觉语言动作基线，也不是终身学习实验。'],
['动作预算','每回合先执行 5 个零动作稳定步，再允许 600 个正常控制动作；20 Hz。抓取前新增的 12 步等待计入这 600 步。'],
['成功判定','完全使用 LIBERO 官方 check_success；首次成功即停止，包括成功动作本身。不额外补拍松手或稳定等待。'],
['录像','两个 256×256 视角拼成 512×256 MP4；最终验证为 20 fps，帧数等于控制步数加 5。'],
['状态与物理','初始化后不直接改活跃环境中的物体 qpos；不改奖励或成功谓词。IK 的私有数据对象不作为环境状态写回。'],
['固定语义','在 rollout 前根据任务 BDDL 构造确定性技能序列；可显式加入几何清障步骤。技能内部有状态反馈，但没有语言模型或语义重规划。']],[3.1,13.9])
head('三类成绩必须分开')
p('开发实验用于暴露和修复问题，可以在保留旧结果的前提下更改代码再跑新版本。正式评测要求冻结实现、保持完整分母并使用未用于开发的状态。当前固定语义验证是在曾反复调试的三个开发状态上进行，因此只能证明该版本覆盖了这些已知初始条件。')
p('修复后 8 个任务从旧固定版本的 0/24 提高到 24/24，是相同任务与开发状态范围的对照。任务 0 和 7 在旧版本各为 3/3，但没有加入最终统一修复运行，不能把两种版本拼成当前版本 10 个任务 100%。同样，固定语义成功不能直接证明 GPT 规划器已达到相同成绩。')
evidence('configs/task_manifest.json；configs/protocol.json；docs/METHOD.md；runs/fixed-repair-validation-v2/run.json')

page(4,'历史运行与失败口径核查')
p('下表统计各运行目录已有 result.json。回合数表示已留存结果的数量，不代表完整计划分母；API 调用数沿用各回合结果字段。中断和环境错误仍保留，排除的旧版本不参与新版本成绩合并。')
table(['运行 ID','回合','成功','API 调用','主要说明'],[
['pilot-gpt6-v1',4,2,32,'2 次重规划预算失败'],
['dev-gpt6-v2',30,2,60,'17 仿真错误、7 重规划失败、4 中断；内存故障批次'],
['dev-gpt6-v3',30,7,182,'21 重规划失败、2 步数超限'],
['dev-gpt6-v4',30,6,177,'24 重规划失败'],
['dev-fixed-v5',30,6,0,'24 次计划执行后未达目标'],
['dev-microwave-fix',3,0,0,'开门改进后关门仍失败'],
['dev-gpt6-v6',30,6,175,'21 重规划失败、3 API 超时'],
['dev-fixed-v7',30,6,0,'原失败八任务合计 0/24'],
['frozen-v1-gpt6',15,12,109,'3 中断；旧冻结实现整批排除'],
['frozen-v2-gpt6',83,33,347,'47 额度错误、3 中断；正式计划 470'],
['fixed-repair-validation-v1',24,23,0,'任务 1 状态 1 清障边界失败'],
['fixed-repair-validation-v2',24,24,0,'最终统一修复版本，全部成功']],[6,1.2,1.2,1.8,6.8])
head('对之前几个关键疑问的核实')
p('开发测试分为调用 GPT 的 dev-gpt6 系列与零调用的 dev-fixed 系列。最近一次 GPT 开发 v6 和旧固定开发 v7 都是任务 0、7 各 3/3，总体 6/30。其余 8 个任务的低成功率在额度耗尽之前已存在，主要问题出在操作实现；v6 另有 3 次请求超时。')
p('正式 v2 实际触及任务 0 和任务 1，并非任务 0 和任务 7。成功的 33 条全部来自任务 0；额度错误包含任务 0 的 14 条和任务 1 的 33 条；任务 1 的状态 36、37、38 中断。早期 v1 因开发中发现任务 9 几何筛选遗漏而作废，未挑选成功记录并入 v2。')
evidence('runs/*/taskXX_stateYYY/result.json；reports/PAUSED_STATUS.json；docs/CHANGELOG.md')

page(5,'共享故障与诊断机制修复')
head('从级联失败追溯首个故障')
p('旧执行链在抓取已失败后仍继续放置，产生大量 grasp_slip 等后续错误，容易误判为多个独立物理问题。新增 debug_fixed.py 遇到第一个技能失败即停止，记录目标、碰撞、接触和当前场景；每次尝试使用独立目录，成功和失败均保留。')
table(['根因','实施修复','作用'],[
['末端位置和姿态坐标系不一致','robosuite 返回 grip-site 位置但原始 wrist-body 四元数；统一为真实 grip-site 世界姿态，并保留 robot0_wrist_quat。','让 IK、控制目标与抓取姿态描述同一个末端。'],
['夹爪闭合轴判断错误','Panda 夹指闭合方向改为 grip-site 局部 x 轴；检查等价 yaw 候选，并收紧 IK 收敛要求。','减少错误夹取方向和不可达候选。'],
['初始物体仍在下落','5 步初始化后再在动作预算内等待 12 步，重新读取物体状态再规划。','避免按过时位置抓取。'],
['薄物下降不足及接触漏判','薄盒定位容差收紧至 3 mm；双指接触覆盖手指网格与指垫，并检查至少 7 cm 的实际抬升。','改善奶酪、黄油和书等物体的真实夹持。'],
['整体包围盒抓取不合适','高物体用上部抓点；摩卡壶抓顶部小柄；空心杯碗使用沿口或杯柄。','避免手掌撞物体、穿过空腔或抓取不稳。']],[3.4,8.2,5.4])
head('零 API 与不可混用版本的保障')
p('固定诊断进程删除继承的 OPENAI_API_KEY，并拦截 socket 连接；不实例化 GPT 规划器。完整配置、预设计划、代码哈希和 source.zip 随运行保存。实现发生变化后必须使用新 run ID，已完成的回合不能覆盖。本轮共登记 91 条修复尝试，全部为零 API 调用。')
p('这些变化修复了坐标变换、技能几何与轨迹，不改变官方成功判定。由于执行实现已经变化，旧 frozen.json 不能用于继续跑混合版本的正式成绩。')
evidence('scripts/debug_fixed.py；docs/FIXED_REPAIR_LOG.md；docs/fixed_repair_attempts.jsonl；tests/')

task_details = [
(1,'奶油奶酪和黄油放入篮子','奶盒遮挡奶酪接近路径，黄油薄盒下降不足；失败抓取继续执行又引出持物错误。','修正薄盒夹持方向和下降容差，在固定计划中显式抓取并移走邻近高物体。第一次统一验证发现状态 1 的奶盒中心距离为 0.120877 m，刚超出 0.12 m 清障范围；扩大到 0.125 m 后该状态需清走两个障碍物，574 步成功。','452 / 574 / 296','这一任务最接近 600 步上限；状态 1 仅余 26 步，成功并不意味着动作预算有充足余量。'),
(2,'打开炉灶并放上摩卡壶','旋钮抓取选中了底部小几何体而非凸起拨片；摩卡壶整体抓取又被壶嘴阻挡。','选取旋钮凸起拨片，绕真实铰链轴拨动并检查官方 Turnon；壶改抓顶部小柄，搬运后检查 Turnon 与 On 同时满足。若开关已经打开则跳过重复拨动。早期首次成功为 216 步，最终统一版本因搬运净空调整为 230 步。','230 / 236 / 233','开关动作和放置动作必须同时满足目标；不能只凭画面中壶接近炉面就判断成功。'),
(3,'碗放入底层抽屉并关闭','抽屉把手选择和方向不匹配；碗可放入后，垂直推关姿态会撞柜体，侧推又受酒架和可达性限制。','依据真实滑动轴选择前方把手，使用碗沿夹取。关闭抽屉采用约 0.45 rad 前倾姿态，预推距离缩短为 5.5 cm，避开上层柜面和侧边酒架。第 9 个任务专项版本在状态 0 达到 333 步成功。','333 / 327 / 348','官方 In 与 Close 均成立；保留了已经放入但无法关抽屉的中间失败记录。'),
(4,'两个杯子分别放到左右盘子','抓取坐标系错误与按杯身整体包围盒抓取叠加，手掌干涉杯沿或不能形成稳定夹持。','共享末端坐标系修复后，采用杯沿夹取，针对各自目标盘计算放置位置与高度，顺序完成两只杯子。','272 / 261 / 263','需区分白杯与黄白杯各自对应的左右盘；最终以两个官方目标分别为真确认。'),
(5,'书放入收纳盒后槽','书本薄体抓取深度不足；按固定 90° 旋转并不保证真实几何长轴与收纳槽一致。','采用薄物精确下降，依据书的主碰撞几何长轴计算槽位对齐姿态，再执行窄槽放置。','133 / 132 / 134','最终验证中动作最少的修复任务；这只是控制步数开销，不等于任何 GPT 正式评测费用结论。'),
(6,'杯子放盘上且布丁放盘右侧','杯体抓取不稳、放置高度不合适，会使第一个子目标失败并影响后续动作。','杯子改用沿口抓取并修正放置高度；布丁使用盒体抓取，将目标设为盘子右侧的官方关系区域。','254 / 257 / 263','同时检查杯子在盘上与布丁在指定右侧区域，避免只核对其中一个子目标。'),
(8,'两个摩卡壶放上炉灶','搬运高度不足，壶与场景或已放置物碰撞；两个物体落点过近可能互相干扰。','将壶搬运净空提高到 0.23 m，使用横向分开的约 ±4 cm 落点；保留顶部小柄抓取，并依据本地 BDDL 满足 Turnon，已经开启则跳过。','297 / 302 / 301','任务名称未完整表达开关要求，验收以本地 BDDL 的全部目标为准。'),
(9,'杯子放入微波炉并关门','从上方进入会撞炉顶；杯沿抓取倾斜后杯柄挂住环境，撤手时把杯子拖出；关门姿态又受到腕部关节与接触滞后影响。','改为抓杯柄，在炉门外降高后从正面送入；依据内部区域和炉底计算放置点，再沿工具轴撤离。门把手抓取增加 90° yaw 以避开腕部极限，关闭目标增加有界 0.06 rad 推进补偿接触柔顺性。','352 / 361 / 353','失败录像保留了杯子先进入又被撤手带出的现象；最终要求杯子 In 与门 Close 同时成立。')]

for n, pair in enumerate([task_details[0:2],task_details[2:4],task_details[4:6],task_details[6:8]],6):
    page(n, f'任务 {pair[0][0]} 与任务 {pair[1][0]} 的故障修复')
    for tid,title,fail,fix,steps,note in pair:
        head(f'任务 {tid}  {title}')
        p('故障表现  '+fail)
        p('实施修复  '+fix)
        p(f'统一验证  状态 0、1、2 全部成功，分别为 {steps} 步。'+note)
        evidence(f'runs/fixed-repair-validation-v2/task{tid:02d}_state000 至 state002；专项迭代见 docs/FIXED_REPAIR_LOG.md')
    if n == 9:
        p('所有关节和接触补偿均通过正常 env.step 动作实现，没有直接修改炉门状态、物体位置或成功阈值。')

page(10,'最终统一验证与录像证据')
p('第二次统一运行 fixed-repair-validation-v2 在同一代码哈希下完成全部 24 回合，成功 24、技能失败 0、API 调用 0。第一轮的 23/24 与第二轮的 24/24 单独归档，未覆盖旧失败回合。')
table(['任务','状态 0 步数','状态 1 步数','状态 2 步数','结果'],[[x[0],*x[4].split(' / '),'3/3'] for x in task_details],[2.2,3.9,3.9,3.9,3.1])
head('验证做到了什么')
p('逐一解码全部 24 个 MP4，核对分辨率 512×256、20 fps、帧数等于动作记录加 5 个稳定步；末帧对应官方成功。任务 1 状态 1 的 574 动作长轨迹另做了一次无 API 独立回放，再次成功，579 帧，最大模拟状态差异为 0。')
p('现有测试共 9 项通过，覆盖协议和关键回归；其中补充末端 grip-site 四元数回归，并使成功动作记录的测试与真实终止行为一致。来源完整性检查通过。此处报告的是先前已经保存的验证结果，本次制作报告未重跑仿真。')
head('证据强度与限制')
p('录像、动作轨迹和官方谓词共同证明这些回合确实通过正常仿真成功。截图只是检索入口，不能独立替代目标判定。由于每个任务只有 3 个用于调试的初始状态，此结果不能证明其他 47 个状态的成功率，也不能据此完成正式统计或声称具有普遍鲁棒性。')
evidence('reports/fixed_repair/summary.json；reports/fixed_repair/replay_task01_state001/replay_result.json；reports/fixed_repair/SUMMARY.md')

page(11,'成功末帧总览')
p('下图为八个修复任务在状态 0 的成功末帧索引，按照任务 1、2、3、4、5、6、8、9 排列。完整视频提供外部相机与腕部相机两个视角；成功时可能仍保持抓取，这是首次官方成功即终止的协议结果。')
d.add_picture(str(ROOT/'reports/fixed_repair/success_frames.jpg'), width=Cm(16.5))
p('图 1  固定语义修复任务成功末帧，完整录像见各回合的 video.mp4', 'Caption')
p('建议阅读顺序为：先看任务失败原因和最终步数，再打开对应视频；如需精确排查某次动作，配合该回合的 skills、action/state 与 scene 文件。所有中间失败尝试仍可按修复日志的 run ID 找到。')
evidence('reports/fixed_repair/success_frames.jpg；reports/fixed_repair/SUMMARY.md')

page(12,'API 用量费用与正式运行暂停')
head('已有用量和费用估计')
table(['指标','核实值','含义'],[
['历史 API 元数据记录','1,076 条','包含失败请求记录，不能称为 1,076 次成功付费请求'],
['已知输入与输出','输入 3,025,494 token\n输出 83,034 token','来自暂停状态保存的累计已知用量'],
['历史累计用量估值','约 32.17 美元','按保存的价格和缓存用量估算；不是账单，未知用量不能当作零'],
['配置预算上限','600 美元','预算保护参数，不是实际支出'],
['正式 470 回合投影','约 82.55 美元','早期 dev-gpt6-v6 校准的无缓存线性推算，不是已发生费用'],
['固定修复及统一验证','0 次 API 调用，0 token','不产生模型 API 用量，仍消耗本地 CPU、渲染和存储资源']],[4,5.2,7.8])
p('费用投影文件采用 2026 年 9 月 7 日保存的每百万 token 单价：输入 10 美元、缓存输入 1 美元、缓存写入 12.5 美元、输出 50 美元。dev-gpt6-v6 的 175 次调用包含 456,083 输入和 14,164 输出 token，按无缓存口径约 5.269 美元；完整正式集的串行时间投影约 5.37 小时。这些都是当时的校准假设，本次没有重新查询价格或核对服务商账单。')
head('暂停原因及尚未修复的问题')
p('正式 v2 收到 HTTP 429，错误信息为 credit_balance_exhausted / insufficient_quota，属于账户额度耗尽。批处理没有在首个额度错误后全局停止，后续累计留下 47 条额度错误。这是调度层待修复的问题，当前固定语义成功并没有修复或验证付费批处理熔断。')
p('因此，不应把 33/470 或 33/83 当作方法最终成功率：前者把大量未执行回合当成失败，后者混入账户与中断问题且任务分布严重不完整。后续如恢复正式评测，应先完成全局额度熔断和新实现冻结，再在独立运行中统计完整分母。当前付费工作仍暂停。')
evidence('reports/PAUSED_STATUS.json；reports/cost_estimate.json；runs/frozen-v2-gpt6/')

page(13,'版本归档与复现说明')
head('仓库与版本标识')
p('本次归档仓库命名为 libero-agentic-manipulation-lab，使用 MEILinnaMelina 账号的私有 GitHub 仓库。版本标签为 v0.1.0-fixed-repair。仓库保存适配代码、依赖锁定文件、原始来源快照、成功和失败实验资料、视频及本 Word 报告；本机 Python 环境、工作临时目录、缓存和凭据文件不进入版本。')
p('仓库地址\nhttps://github.com/MEILinnaMelina/libero-agentic-manipulation-lab')
table(['标识','值'],[
['最终实验代码 SHA-256','29ccbc6c9f3b91351ce83c7fb0e225cd4dc394c8c74dde093df5aea80885e803'],
['LIBERO 来源提交','8f1084e3132a39270c3a13ebe37270a43ece2a01'],
['RoboEval 来源提交','988687694c54b8f0e4038fee8c9dade2621b5743']],[4,13])
p('实验代码哈希用于核对最终统一实验的执行源码；Git 提交还包含报告与归档元数据，二者不是同一个标识。报告不写入其自身所属提交的哈希，以避免自引用；可通过上述标签和 Git 历史定位整个交付版本。')
head('固定语义复现入口')
p('在 README 指引下创建 Python 3.8.13 环境、安装锁定依赖、应用 Windows 补丁并验证来源。已有录像和结果可以直接读取；需要重新执行时，使用新的 run ID，避免覆盖历史证据。以下为固定语义入口，不涉及 API：')
p('.\\.venv-sim\\python.exe scripts\\debug_fixed.py --tasks 1,2,3,4,5,6,8,9 --states 0,1,2 --run-id fixed-repair-new')
p('已有结果的源代码快照位于 runs/fixed-repair-validation-v2/source.zip。回放可用 scripts/replay.py 读取指定回合的动作文件。当前环境路径是 Windows 本机路径，其他机器需按 README 重建环境；vendor 快照包含较多仿真资产，首次克隆会占用较多下载时间和磁盘。')
p('LIBERO 的 MIT 与 RoboEval 的 Apache-2.0 许可证及来源说明一并保留。Git 归档保留文件字节，避免换行自动转换影响来源校验和已冻结代码哈希。')
evidence('README.md；.gitattributes；docs/provenance/sources.json；runs/fixed-repair-validation-v2/source.zip')

page(14,'未完成事项与后续执行条件')
p('当前阶段达成的是用户后续指定的八任务固定语义修复目标；原始评测清单中的正式实验目标仍有缺口。本报告不把准备好的脚本等同于已经执行并验证的研究结论。')
table(['事项','目前状态','后续应采取的动作'],[
['全十任务同版本开发回归','最终修复统一验证只覆盖 8 个任务','将任务 0、7 加入相同修复版本，检查共享技能修改是否导致回归。'],
['未见初始状态的固定语义验证','未建立修复版本的独立正式成绩','如继续固定实验，使用尚未用于调试的状态，并保持固定版本和完整计划分母。'],
['GPT 正式评测','旧版因额度耗尽暂停；新修复版未运行 GPT','保持暂停，恢复前先解决额度熔断并重新校准和冻结；不混合旧结果。'],
['固定与无重规划匹配消融','两个正式条件各计划 50 回合，均未完成执行','在同一冻结版本和匹配状态下独立执行，避免不同实现之间伪对照。'],
['正式统计与中英文研究报告','尚缺完整实验数据','数据完整后再给出置信区间、方法比较与正式结论；本文件为阶段工作报告。'],
['鲁棒性与预算余量','任务 1 状态 1 已用 574/600 步','优先优化冗余清障和路径，再用未调试状态检验，不能通过扩大动作预算掩盖问题。']],[4,5.4,7.6])
head('本轮工作带来的认识')
p('固定语义原先的失败不是“没有使用更强语言模型”这一单一原因。末端坐标系、夹爪轴向、碰撞几何选择、薄物接近精度和狭窄空间路径等底层问题，会使正确的高层任务顺序也无法完成。先保留首个失败、核对物理接触，再逐项修复，比将所有失败概括成模型规划错误更符合现有证据。')
p('另一方面，调试状态上的全部成功仍需要独立验证。应继续保持失败记录、清晰分母、冻结版本和真实官方成功判定，才能区分实现修复效果、初始状态难度和规划策略贡献。以上后续项目只是交接建议，本次报告归档没有自动启动这些实验。')

page(15,'证据索引与交付文件')
p('以下路径均相对于仓库根目录。JSON 和逐回合目录用于审计具体数值，Markdown 用于快速阅读，MP4 用于观察操作过程；三者共同构成本报告的证据链。')
table(['内容','路径'],[
['本 Word 工作报告','reports/work_report/LIBERO_工作报告_2026-09-07至08.docx'],
['目标及完成状态','TASK_LIST.md；docs/TASK_LIST.original.md'],
['方法和版本变更','docs/METHOD.md；docs/CHANGELOG.md'],
['来源及平台适配','docs/provenance/sources.json；windows_patches.json'],
['协议和状态划分','configs/protocol.json；configs/task_manifest.json'],
['API 暂停与费用','reports/PAUSED_STATUS.json；reports/cost_estimate.json'],
['正式运行原始证据','runs/frozen-v1-gpt6/；runs/frozen-v2-gpt6/'],
['逐项修复过程','docs/FIXED_REPAIR_LOG.md'],
['全部修复尝试索引','docs/fixed_repair_attempts.jsonl'],
['修复结果和视频导航','reports/fixed_repair/SUMMARY.md；summary.json'],
['最终统一运行与源码','runs/fixed-repair-validation-v2/run.json；source.zip'],
['24 回合动作和视频','runs/fixed-repair-validation-v2/taskXX_stateYYY/'],
['精确回放结果','reports/fixed_repair/replay_task01_state001/replay_result.json'],
['物理一致性验证','reports/memory_equivalence.json'],
['复现指引及依赖','README.md；configs/requirements-sim.lock']],[5.2,11.8])
p('录像目录命名示例：task01_state001 表示任务 1、官方初始状态索引 1。该目录中的 video.mp4 为 574 动作成功回合的原始录像。其他回合按同一规则定位。')
p('报告中费用、时间投影和运行结果均来自本地历史材料，不包含对当前 API 价格或账户余额的新查询。最终交付的远端链接、Git 提交和版本标签在交付消息中核对，以仓库实际状态为准。')
d.save(OUT)
print(OUT)
